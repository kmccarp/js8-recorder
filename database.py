"""
SQLite database layer for JS8 Recorder.
"""

import sqlite3
from pathlib import Path
from datetime import datetime
from typing import Optional

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None


def format_snr(value) -> str:
    """Format SNR to JS8Call display format. '2' -> '+02', '-10' -> '-10'"""
    if value is None or value == "":
        return ""
    try:
        num = int(str(value).replace("+", ""))
        return f"{num:+03d}"
    except ValueError:
        return str(value)  # Return as-is if can't parse (old data)


def grid_to_latlon(grid: str) -> tuple:
    """Convert 4-character Maidenhead grid to center lat/lon coordinates.

    Returns (lat, lon) tuple or None if invalid grid.
    """
    if not grid or len(grid) < 4:
        return None
    try:
        # A = -180 lon, A = -90 lat
        # Each field = 20 lon x 10 lat degrees
        # Each square = 2 lon x 1 lat degrees
        lon = (ord(grid[0].upper()) - ord('A')) * 20 - 180 + int(grid[2]) * 2 + 1
        lat = (ord(grid[1].upper()) - ord('A')) * 10 - 90 + int(grid[3]) * 1 + 0.5
        return (lat, lon)
    except (ValueError, IndexError):
        return None


def get_adjacent_grids(grid: str) -> list:
    """Return list of 8 adjacent grid squares for a Maidenhead grid."""
    if len(grid) < 4:
        return []

    try:
        field_lon = ord(grid[0].upper())  # A-R
        field_lat = ord(grid[1].upper())  # A-R
        sq_lon = int(grid[2])             # 0-9
        sq_lat = int(grid[3])             # 0-9
    except (ValueError, IndexError):
        return []

    adjacent = []
    for d_lon in [-1, 0, 1]:
        for d_lat in [-1, 0, 1]:
            if d_lon == 0 and d_lat == 0:
                continue  # Skip center (the input grid itself)

            new_sq_lon = sq_lon + d_lon
            new_sq_lat = sq_lat + d_lat
            new_field_lon = field_lon
            new_field_lat = field_lat

            # Handle wraparound
            if new_sq_lon < 0:
                new_sq_lon = 9
                new_field_lon -= 1
            elif new_sq_lon > 9:
                new_sq_lon = 0
                new_field_lon += 1

            if new_sq_lat < 0:
                new_sq_lat = 9
                new_field_lat -= 1
            elif new_sq_lat > 9:
                new_sq_lat = 0
                new_field_lat += 1

            # Bounds check (A-R)
            if new_field_lon < ord('A') or new_field_lon > ord('R'):
                continue
            if new_field_lat < ord('A') or new_field_lat > ord('R'):
                continue

            new_grid = f"{chr(new_field_lon)}{chr(new_field_lat)}{new_sq_lon}{new_sq_lat}"
            adjacent.append(new_grid)

    return adjacent


def format_age(timestamp: str) -> str:
    """Format timestamp as human-readable age (e.g., '2h ago', '1d ago')."""
    if not timestamp:
        return ""
    try:
        dt = datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S")
        now = datetime.utcnow()
        delta = now - dt

        minutes = int(delta.total_seconds() / 60)
        hours = int(delta.total_seconds() / 3600)
        days = delta.days
        weeks = days // 7

        if minutes < 60:
            return f"{minutes}m ago"
        elif hours < 24:
            return f"{hours}h ago"
        elif days < 7:
            return f"{days}d ago"
        else:
            return f"{weeks}w ago"
    except ValueError:
        return ""


class Database:
    def __init__(self, db_path: str = "js8_log.db"):
        self.db_path = Path(db_path)
        self.conn = sqlite3.connect(str(self.db_path), check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self._create_tables()

    def _create_tables(self):
        """Create database tables if they don't exist."""
        cursor = self.conn.cursor()

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS directed_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                callsign TEXT NOT NULL,
                timestamp TEXT NOT NULL,
                my_snr_of_them TEXT,
                their_snr_of_me TEXT,
                message TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS callsign_grids (
                callsign TEXT PRIMARY KEY,
                grid TEXT NOT NULL,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT
            )
        """)

        # Migration: Add band column if it doesn't exist
        cursor.execute("PRAGMA table_info(directed_messages)")
        columns = [col[1] for col in cursor.fetchall()]
        if "band" not in columns:
            cursor.execute("ALTER TABLE directed_messages ADD COLUMN band TEXT")

        self.conn.commit()

    def add_message(self, callsign: str, timestamp: str, my_snr_of_them: str,
                    their_snr_of_me: str, message: str, band: str = "") -> int:
        """Add a directed message to the database. Returns the row ID."""
        cursor = self.conn.cursor()
        cursor.execute("""
            INSERT INTO directed_messages (callsign, timestamp, my_snr_of_them, their_snr_of_me, message, band)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (callsign, timestamp, my_snr_of_them, their_snr_of_me, message, band))
        self.conn.commit()
        return cursor.lastrowid

    def add_grid(self, callsign: str, grid: str):
        """Add or update a callsign-grid mapping. Won't overwrite existing grid with empty."""
        cursor = self.conn.cursor()
        now = datetime.utcnow().isoformat()
        if grid:
            # Insert or update with the new grid
            cursor.execute("""
                INSERT INTO callsign_grids (callsign, grid, updated_at)
                VALUES (?, ?, ?)
                ON CONFLICT(callsign) DO UPDATE SET grid = ?, updated_at = ?
            """, (callsign, grid, now, grid, now))
        else:
            # Only insert if callsign doesn't exist, don't overwrite existing grid
            cursor.execute("""
                INSERT OR IGNORE INTO callsign_grids (callsign, grid, updated_at)
                VALUES (?, '', ?)
            """, (callsign, now))
        self.conn.commit()

    def update_grid(self, callsign: str, grid: str):
        """Update grid for a callsign (allows setting to empty)."""
        cursor = self.conn.cursor()
        cursor.execute("""
            UPDATE callsign_grids SET grid = ?, updated_at = ?
            WHERE callsign = ?
        """, (grid, datetime.utcnow().isoformat(), callsign))
        self.conn.commit()

    def get_all_messages(self) -> list:
        """Get all directed messages, oldest first (for display with newest at top)."""
        cursor = self.conn.cursor()
        cursor.execute("""
            SELECT id, callsign, timestamp, my_snr_of_them, their_snr_of_me, message, band
            FROM directed_messages
            ORDER BY timestamp ASC
        """)
        return [dict(row) for row in cursor.fetchall()]

    def get_all_grids(self) -> list:
        """Get all callsign-grid mappings."""
        cursor = self.conn.cursor()
        cursor.execute("""
            SELECT callsign, grid
            FROM callsign_grids
            ORDER BY callsign
        """)
        return [dict(row) for row in cursor.fetchall()]

    def get_grids_with_snr_stats(self) -> list:
        """Get all callsign-grid mappings with SNR statistics and bands."""
        cursor = self.conn.cursor()
        cursor.execute("""
            SELECT
                cg.callsign,
                cg.grid,
                MAX(CAST(dm.my_snr_of_them AS INTEGER)) as max_my_snr,
                MIN(CAST(dm.my_snr_of_them AS INTEGER)) as min_my_snr,
                MAX(CAST(dm.their_snr_of_me AS INTEGER)) as max_their_snr,
                MIN(CAST(dm.their_snr_of_me AS INTEGER)) as min_their_snr,
                MAX(dm.timestamp) as last_contact,
                COUNT(dm.id) as contact_count,
                GROUP_CONCAT(DISTINCT dm.band) as bands
            FROM callsign_grids cg
            LEFT JOIN directed_messages dm ON cg.callsign = dm.callsign
            GROUP BY cg.callsign, cg.grid
            ORDER BY cg.callsign
        """)
        return [dict(row) for row in cursor.fetchall()]

    def lookup_by_grid(self, grid: str) -> list:
        """Find callsigns by grid square, sorted by likelihood to hear you."""
        cursor = self.conn.cursor()
        cursor.execute("""
            SELECT
                cg.callsign,
                cg.grid,
                AVG(CAST(dm.their_snr_of_me AS INTEGER)) as avg_their_snr,
                MAX(CAST(dm.their_snr_of_me AS INTEGER)) as max_their_snr,
                COUNT(dm.id) as contact_count,
                MAX(dm.timestamp) as last_contact
            FROM callsign_grids cg
            INNER JOIN directed_messages dm ON cg.callsign = dm.callsign
            WHERE cg.grid = ?
            GROUP BY cg.callsign, cg.grid
            ORDER BY avg_their_snr DESC
        """, (grid.upper(),))
        return [dict(row) for row in cursor.fetchall()]

    def get_message_count(self) -> int:
        """Get total message count."""
        cursor = self.conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM directed_messages")
        return cursor.fetchone()[0]

    def delete_message(self, message_id: int):
        """Delete a single message by ID."""
        cursor = self.conn.cursor()
        cursor.execute("DELETE FROM directed_messages WHERE id = ?", (message_id,))
        self.conn.commit()

    def delete_messages(self, message_ids: list):
        """Delete multiple messages by ID (batch delete)."""
        if not message_ids:
            return
        cursor = self.conn.cursor()
        placeholders = ",".join("?" * len(message_ids))
        cursor.execute(f"DELETE FROM directed_messages WHERE id IN ({placeholders})", message_ids)
        self.conn.commit()

    def delete_callsign_grid(self, callsign: str):
        """Delete just the callsign-grid mapping."""
        cursor = self.conn.cursor()
        cursor.execute("DELETE FROM callsign_grids WHERE callsign = ?", (callsign,))
        self.conn.commit()

    def delete_callsign_with_messages(self, callsign: str):
        """Delete callsign grid entry AND all messages for that callsign."""
        cursor = self.conn.cursor()
        cursor.execute("DELETE FROM directed_messages WHERE callsign = ?", (callsign,))
        cursor.execute("DELETE FROM callsign_grids WHERE callsign = ?", (callsign,))
        self.conn.commit()

    def get_message_count_for_callsign(self, callsign: str) -> int:
        """Get count of messages for a callsign."""
        cursor = self.conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM directed_messages WHERE callsign = ?", (callsign,))
        return cursor.fetchone()[0]

    def get_setting(self, key: str, default: str = "") -> str:
        """Get a setting value."""
        cursor = self.conn.cursor()
        cursor.execute("SELECT value FROM settings WHERE key = ?", (key,))
        row = cursor.fetchone()
        return row["value"] if row else default

    def set_setting(self, key: str, value: str):
        """Set a setting value."""
        cursor = self.conn.cursor()
        cursor.execute("""
            INSERT INTO settings (key, value)
            VALUES (?, ?)
            ON CONFLICT(key) DO UPDATE SET value = ?
        """, (key, value, value))
        self.conn.commit()

    def export_to_excel(self, output_path: str) -> bool:
        """Export database to Excel file. Returns True on success."""
        if Workbook is None:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")

        wb = Workbook()

        # Sheet 1: Directed Messages
        ws1 = wb.active
        ws1.title = "Directed Messages"
        ws1.append(["Callsign", "QRZ", "Timestamp (UTC)", "Band", "My SNR of Them", "Their SNR of Me", "Message"])

        for msg in self.get_all_messages():
            callsign = msg["callsign"]
            qrz_url = f"https://www.qrz.com/db/{callsign}" if callsign else ""
            ws1.append([
                callsign,
                qrz_url,
                msg["timestamp"],
                msg.get("band", ""),
                format_snr(msg["my_snr_of_them"]),
                format_snr(msg["their_snr_of_me"]),
                msg["message"]
            ])

        ws1.column_dimensions["A"].width = 12
        ws1.column_dimensions["B"].width = 35
        ws1.column_dimensions["C"].width = 20
        ws1.column_dimensions["D"].width = 8
        ws1.column_dimensions["E"].width = 15
        ws1.column_dimensions["F"].width = 15
        ws1.column_dimensions["G"].width = 50

        # Sheet 2: Callsign Grids with SNR stats
        ws2 = wb.create_sheet("Callsign Grids")
        ws2.append(["Callsign", "QRZ", "Grid", "Bands", "Max My SNR", "Min My SNR",
                    "Max Their SNR", "Min Their SNR", "Last Contact"])

        for entry in self.get_grids_with_snr_stats():
            callsign = entry["callsign"]
            qrz_url = f"https://www.qrz.com/db/{callsign}" if callsign else ""
            # Format bands - filter out empty strings and join
            bands = entry.get("bands", "") or ""
            bands_list = [b for b in bands.split(",") if b]
            bands_formatted = ", ".join(sorted(set(bands_list)))
            ws2.append([
                callsign,
                qrz_url,
                entry["grid"],
                bands_formatted,
                format_snr(entry["max_my_snr"]),
                format_snr(entry["min_my_snr"]),
                format_snr(entry["max_their_snr"]),
                format_snr(entry["min_their_snr"]),
                format_age(entry["last_contact"])
            ])

        ws2.column_dimensions["A"].width = 12
        ws2.column_dimensions["B"].width = 35
        ws2.column_dimensions["C"].width = 8
        ws2.column_dimensions["D"].width = 15
        ws2.column_dimensions["E"].width = 12
        ws2.column_dimensions["F"].width = 12
        ws2.column_dimensions["G"].width = 14
        ws2.column_dimensions["H"].width = 14
        ws2.column_dimensions["I"].width = 12

        wb.save(output_path)
        return True

    def close(self):
        """Close the database connection."""
        self.conn.close()
